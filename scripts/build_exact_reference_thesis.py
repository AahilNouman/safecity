# -*- coding: utf-8 -*-
"""
SafeCity: SMART COMMUNITY-BASED WOMEN PROTECTION AND INCIDENT REPORTING PLATFORM
Exact Word Document Generator based on the Reference Thesis Structure.
Presidency University, Bengaluru, India
Department of Computer Science and Engineering, School of Computer Science and Engineering
Supervisor: Mr. Muthuraju V., Assistant Professor

Authors:
- Mohammed Aahil Nouman (USN: 20231CSE0663) - Project Lead
- Akshat Singh (USN: 20231CSE0609)
- Siddiq (USN: 20231CSE0475)
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_chapter_heading(doc, chapter_num, title):
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(24)
    p1.paragraph_format.space_after = Pt(4)
    p1.paragraph_format.keep_with_next = True
    run1 = p1.add_run(chapter_num)
    run1.font.name = 'Times New Roman'
    run1.font.size = Pt(16)
    run1.font.bold = True
    run1.font.color.rgb = RGBColor(15, 23, 42)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(18)
    p2.paragraph_format.keep_with_next = True
    run2 = p2.add_run(title)
    run2.font.name = 'Times New Roman'
    run2.font.size = Pt(16)
    run2.font.bold = True
    run2.font.color.rgb = RGBColor(15, 23, 42)

def add_subtopic_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)

def add_sub_subtopic_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59)

def add_body_paragraph(doc, text, bold_prefix=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)

    r_body = p.add_run(text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(12)
    r_body.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet_point(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)

    r_body = p.add_run(text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(12)
    r_body.font.color.rgb = RGBColor(30, 41, 59)

def add_callout_box(doc, title, items, border_color="0284C7", bg_color="F0F9FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(4)
    r0 = p0.add_run(title)
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(12)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(15, 23, 42)

    for it_pre, it_txt in items:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.3
        rp = p.add_run(f"• {it_pre}: ")
        rp.font.name = 'Times New Roman'
        rp.font.size = Pt(11)
        rp.font.bold = True
        rp.font.color.rgb = RGBColor(15, 23, 42)
        rt = p.add_run(it_txt)
        rt.font.name = 'Times New Roman'
        rt.font.size = Pt(11)
        rt.font.color.rgb = RGBColor(51, 65, 85)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(0)
    p_post.paragraph_format.space_after = Pt(4)

def add_figure_with_caption(doc, img_path, fig_num_str, caption_text, width_inches=5.6):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(14)
        r_fig = p_cap.add_run(fig_num_str + ": ")
        r_fig.font.name = 'Times New Roman'
        r_fig.font.size = Pt(10)
        r_fig.font.bold = True
        r_fig.font.color.rgb = RGBColor(15, 23, 42)

        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(51, 65, 85)
    else:
        print(f"Warning: Figure image path not found: {img_path}")

def add_table_caption(doc, tbl_num_str, caption_text):
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(10)
    p_cap.paragraph_format.space_after = Pt(4)
    r_tbl = p_cap.add_run(tbl_num_str + ": ")
    r_tbl.font.name = 'Times New Roman'
    r_tbl.font.size = Pt(10)
    r_tbl.font.bold = True
    r_tbl.font.color.rgb = RGBColor(15, 23, 42)

    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(10)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(51, 65, 85)

print("Exact reference tools initialized successfully.")
