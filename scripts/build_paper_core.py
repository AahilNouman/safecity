# build_complete_safecity_ieee_paper.py
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()

# 0.75-inch margins for standard conference layout
for s in doc.sections:
    s.top_margin = Inches(0.75)
    s.bottom_margin = Inches(0.75)
    s.left_margin = Inches(0.75)
    s.right_margin = Inches(0.75)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(17)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)

def add_authors(names, affiliations, supervisor):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(names + '\n')
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.bold = True

    run_aff = p.add_run(affiliations + '\n')
    run_aff.font.name = 'Times New Roman'
    run_aff.font.size = Pt(9.5)
    run_aff.font.italic = True

    run_sup = p.add_run(supervisor)
    run_sup.font.name = 'Times New Roman'
    run_sup.font.size = Pt(9.5)
    run_sup.font.bold = True
    p.paragraph_format.space_after = Pt(14)

def add_abstract_keywords(abstract_text, keywords_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08

    r_abs_lbl = p.add_run('Abstract — ')
    r_abs_lbl.font.name = 'Times New Roman'
    r_abs_lbl.font.size = Pt(9.5)
    r_abs_lbl.font.bold = True

    r_abs = p.add_run(abstract_text)
    r_abs.font.name = 'Times New Roman'
    r_abs.font.size = Pt(9.5)

    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw.paragraph_format.space_after = Pt(14)

    r_kw_lbl = p_kw.add_run('Keywords — ')
    r_kw_lbl.font.name = 'Times New Roman'
    r_kw_lbl.font.size = Pt(9.5)
    r_kw_lbl.font.bold = True

    r_kw = p_kw.add_run(keywords_text)
    r_kw.font.name = 'Times New Roman'
    r_kw.font.size = Pt(9.5)

def add_heading1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)

def add_heading2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(30, 41, 59)

def add_body(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(30, 41, 59)

def add_bullet(lead, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    r_lead = p.add_run(lead + ' ')
    r_lead.font.name = 'Times New Roman'
    r_lead.font.size = Pt(10)
    r_lead.font.bold = True
    r_text = p.add_run(text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(10)

def add_code_listing(caption, code_text):
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(8)
    p_cap.paragraph_format.space_after = Pt(3)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(9)
    r_cap.font.italic = True

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, 'F8FAFC')
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Setup completed successfully.")
