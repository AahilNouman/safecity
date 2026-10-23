# -*- coding: utf-8 -*-
"""
Master Compiler for SafeCity Exact Reference B.Tech Thesis Report
Strictly adhering to Presidency University Guidelines and Reference Template.
Authors:
- Mohammed Aahil Nouman (USN: 20231CSE0663) - Project Lead
- Akshat Singh (USN: 20231CSE0609)
- Siddiq (USN: 20231CSE0475)
Supervisor: Mr. Muthuraju V., Assistant Professor, Dept. of CSE
"""

import os
import sys
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

import build_exact_reference_thesis as helper
from exact_ref_part1 import build_front_matter
from exact_ref_part2 import build_chapter_1, build_chapter_2, build_chapter_3, build_chapter_4
from exact_ref_part3 import (
    build_chapter_5, build_chapter_6, build_chapter_7,
    build_chapter_8, build_chapter_9, build_conclusion_future_work, build_references
)
from exact_ref_part4 import (
    build_appendix_a, build_appendix_b, build_appendix_c
)

def build_entire_thesis():
    print("Initializing Document with Presidency University Margins & Typography...")
    doc = Document()

    # Configure A4 page size and margins
    for sec in doc.sections:
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.25)  # 1.25" Binding Gutter
        sec.right_margin = Inches(1.0)

        # Header
        header = sec.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Dept. of Computer Science & Engineering, Presidency University, Bengaluru")
        r_hdr.font.name = "Times New Roman"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.italic = True
        r_hdr.font.color.rgb = RGBColor(148, 163, 184)

        # Footer
        footer = sec.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_ftr = p_ftr.add_run("SafeCity: Smart Community-Based Women Protection Platform")
        r_ftr.font.name = "Times New Roman"
        r_ftr.font.size = Pt(8.5)
        r_ftr.font.italic = True
        r_ftr.font.color.rgb = RGBColor(148, 163, 184)

    # Set Normal Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)

    print("Building Part 1: Front Matter...")
    build_front_matter(doc, helper)

    print("Building Part 2: Chapters 1 to 4...")
    build_chapter_1(doc, helper)
    build_chapter_2(doc, helper)
    build_chapter_3(doc, helper)
    build_chapter_4(doc, helper)

    print("Building Part 3: Chapters 5 to 9, Conclusion & References...")
    build_chapter_5(doc, helper)
    build_chapter_6(doc, helper)
    build_chapter_7(doc, helper)
    build_chapter_8(doc, helper)
    build_chapter_9(doc, helper)
    build_conclusion_future_work(doc, helper)
    build_references(doc, helper)

    print("Building Part 4: Appendices A, B (Screenshots), and C (Enclosures)...")
    build_appendix_a(doc, helper)
    build_appendix_b(doc, helper)
    build_appendix_c(doc, helper)

    brain_docx = r"C:\Users\AAHIL\.gemini\antigravity\brain\486918fb-9a46-4b06-97ff-51e77234ab28\SafeCity_Complete_BTech_Thesis_Report.docx"
    doc.save(brain_docx)
    print(f"Document saved successfully at: {brain_docx}")

    # Copy to workspace
    ws_docx = r"C:\Users\AAHIL\OneDrive\Pictures\DOCUMENTS\astrasafe-fullstack\SafeCity_Complete_BTech_Thesis_Report.docx"
    try:
        shutil.copy2(brain_docx, ws_docx)
        print(f"Copied to workspace at: {ws_docx}")
    except Exception as e:
        print(f"Workspace copy note: {e}")

    return brain_docx

if __name__ == "__main__":
    build_entire_thesis()
