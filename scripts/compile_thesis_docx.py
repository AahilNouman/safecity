# -*- coding: utf-8 -*-
"""
Master compiler script for the SafeCity Complete B.Tech Capstone Thesis.
Presidency University, Bengaluru
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Import helper functions and builder definitions
from build_full_thesis_report import (
    set_cell_background,
    set_cell_margins,
    set_table_borders,
    add_chapter_heading,
    add_subtopic_heading,
    add_sub_subtopic_heading,
    add_body_paragraph,
    add_bullet_point,
    add_callout_box,
    add_figure_with_caption,
    add_table_caption
)

from part1_front_matter import part1_front_matter
from part2_chapters_1_to_3 import part2_text
from part3_chapters_4_to_6 import part3_text
from part4_chapters_7_to_9 import part4_text
from part5_conclusion_references_appendices import part5_text

# Execute text definitions in global namespace
namespace = {
    'doc': None,
    'set_cell_background': set_cell_background,
    'set_cell_margins': set_cell_margins,
    'set_table_borders': set_table_borders,
    'add_chapter_heading': add_chapter_heading,
    'add_subtopic_heading': add_subtopic_heading,
    'add_sub_subtopic_heading': add_sub_subtopic_heading,
    'add_body_paragraph': add_body_paragraph,
    'add_bullet_point': add_bullet_point,
    'add_callout_box': add_callout_box,
    'add_figure_with_caption': add_figure_with_caption,
    'add_table_caption': add_table_caption,
    'Inches': Inches,
    'Pt': Pt,
    'RGBColor': RGBColor,
    'WD_ALIGN_PARAGRAPH': WD_ALIGN_PARAGRAPH,
    'WD_TABLE_ALIGNMENT': WD_TABLE_ALIGNMENT
}

exec(part1_front_matter, namespace)
exec(part2_text, namespace)
exec(part3_text, namespace)
exec(part4_text, namespace)
exec(part5_text, namespace)

def main():
    print("Initializing Document...")
    doc = Document()

    # Configure A4 page size and thesis margins
    for sec in doc.sections:
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.25)  # 1.25" Binding Gutter
        sec.right_margin = Inches(1.0)

        # Header
        header = sec.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Dept. of Computer Science & Engineering, Presidency University, Bengaluru")
        r_hdr.font.name = "Times New Roman"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.italic = True
        r_hdr.font.color.rgb = RGBColor(148, 163, 184)

        # Footer
        footer = sec.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_ftr = p_ftr.add_run("SafeCity: Smart Community-Based Women Protection Platform")
        r_ftr.font.name = "Times New Roman"
        r_ftr.font.size = Pt(8.5)
        r_ftr.font.italic = True
        r_ftr.font.color.rgb = RGBColor(148, 163, 184)

    # Set Normal Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)

    print("Building Front Matter...")
    namespace['build_front_matter'](doc)

    print("Building Chapters 1 to 3...")
    namespace['build_chapter_1'](doc)
    namespace['build_chapter_2'](doc)
    namespace['build_chapter_3'](doc)

    print("Building Chapters 4 to 6...")
    namespace['build_chapter_4'](doc)
    namespace['build_chapter_5'](doc)
    namespace['build_chapter_6'](doc)

    print("Building Chapters 7 to 9...")
    namespace['build_chapter_7'](doc)
    namespace['build_chapter_8'](doc)
    namespace['build_chapter_9'](doc)

    print("Building Conclusion, References, and Appendices...")
    namespace['build_conclusion_future_work'](doc)
    namespace['build_references'](doc)
    namespace['build_appendix_a'](doc)
    namespace['build_appendix_b'](doc)
    namespace['build_appendix_c'](doc)

    out_docx = r"C:\Users\AAHIL\OneDrive\Pictures\DOCUMENTS\astrasafe-fullstack\SafeCity_Complete_BTech_Thesis_Report.docx"
    doc.save(out_docx)
    print(f"SafeCity Complete B.Tech Thesis DOCX generated at: {out_docx}")

if __name__ == "__main__":
    main()
